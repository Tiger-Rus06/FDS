"""Builds the combined course-end project report (Part A Titanic, Part B Bank Marketing) from report_template.docx.

Run Data_Science_Case_Studies.ipynb first (it creates figures/ and model/), then:
    python make_report.py
This writes Data_Science_Case_Studies_Report.docx. Open it in Word and press Ctrl+A, F9 to refresh page
numbers (export_pdf.ps1 does this automatically and also saves the PDF).

Fill in DETAILS below. Values still in <angle brackets> are printed in red as a reminder.
"""
import copy
import json
import re
from pathlib import Path

import joblib
import matplotlib
import pandas as pd

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DETAILS = {
    "title": "Data Science Case Studies: Titanic Survival Prediction and Bank Marketing Classification",
    "course": "Fundamentals of Data Science (22CSC59)",
    "students": ["Sarapu Rushil (160124733197)"],
    "section": "CSE-3",
    "year": "III Year / V Semester",
    "faculty": "S. Durga Devi",
    "designation": "Assistant Professor",
    "degree": "BACHELOR OF ENGINEERING",
    "branch": "COMPUTER SCIENCE AND ENGINEERING",
    "department": "Department of Computer Science and Engineering,",
    "academic_year": "[2026-2027]",
}

BASE = Path(__file__).parent
FIG = BASE / "figures"
OUT = BASE / "Data_Science_Case_Studies_Report.docx"
FONT = "Times New Roman"
RED = RGBColor(0xC0, 0x00, 0x00)

info = json.loads((BASE / "model" / "titanic_metrics.json").read_text())
bank = json.loads((BASE / "model" / "bank_metrics.json").read_text())
M = info["metrics"]
(TN, FP), (FN, TP) = info["confusion_matrix"]
df = pd.read_csv(BASE / "data" / "titanic.csv")


# ---------------------------------------------------------------- diagrams
def make_diagrams():
    (FIG / "report").mkdir(exist_ok=True)
    def box_chart(steps, path, title, colors_):
        fig, ax = plt.subplots(figsize=(1.75 * len(steps), 2.6))
        ax.set_xlim(0, len(steps) * 2.6)
        ax.set_ylim(0, 2.6)
        ax.axis("off")
        for i, (head, sub) in enumerate(steps):
            x = i * 2.6 + 0.1
            ax.add_patch(FancyBboxPatch((x, 0.5), 2.2, 1.4, boxstyle="round,pad=0.05",
                                        fc=colors_[i % len(colors_)], ec="#333333", lw=1))
            ax.text(x + 1.1, 1.45, head, ha="center", va="center", fontsize=9.5, weight="bold", family="serif")
            ax.text(x + 1.1, 0.9, sub, ha="center", va="center", fontsize=8, family="serif")
            if i < len(steps) - 1:
                ax.annotate("", xy=(x + 2.55, 1.2), xytext=(x + 2.3, 1.2),
                            arrowprops=dict(arrowstyle="->", lw=1.5, color="#333333"))
        ax.set_title(title, fontsize=12, family="serif", weight="bold")
        fig.savefig(path, bbox_inches="tight", dpi=200)
        plt.close(fig)

    box_chart([("Data", "titanic.csv\n891 rows"), ("EDA", "statistics,\nvisualisation"),
               ("Preprocessing", "imputation,\nencoding, scaling"), ("Split", "80% train\n20% test"),
               ("Training", "logistic\nregression"), ("Evaluation", "metrics,\nROC curve"),
               ("Interpretation", "odds ratios,\np-values"), ("Deployment", "Streamlit\napp")],
              FIG / "report" / "workflow.png", "Project Workflow",
              ["#dbe9f6", "#e3f1e0", "#fdebd3", "#f3e0f0"])


# ---------------------------------------------------------------- low-level helpers
_bm_id = [100]


def set_style_font(style, size=None, bold=None, color=None):
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in list(rfonts.attrib):
        if "Theme" in a or "theme" in a:
            del rfonts.attrib[a]
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), FONT)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    style.font.color.rgb = color or RGBColor(0, 0, 0)


def fmt_run(run, size=12, bold=False, italic=False, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_bookmark(p, name):
    _bm_id[0] += 1
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(_bm_id[0]))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(_bm_id[0]))
    ppr = p._p.pPr
    (ppr.addnext if ppr is not None else lambda e: p._p.insert(0, e))(start)
    p._p.append(end)


def add_field(p, instr, placeholder="1", size=12, bold=False):
    def r(child):
        run = p.add_run()
        fmt_run(run, size, bold)
        run._r.append(child)
        return run
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); r(fc)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r(it)
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); r(fc)
    run = p.add_run(placeholder); fmt_run(run, size, bold)
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r(fc)


def spacing(p, before=0, after=6, line=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    p.alignment = align


def rich(p, text, size=12, color=None):
    """Add text to p, rendering **bold** and *italic* segments."""
    for part in re.split(r"(\*\*.+?\*\*|\*.+?\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            fmt_run(p.add_run(part[2:-2]), size, bold=True, color=color)
        elif part.startswith("*"):
            fmt_run(p.add_run(part[1:-1]), size, italic=True, color=color)
        else:
            fmt_run(p.add_run(part), size, color=color)


def shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def set_text(p, text):
    """Replace a template paragraph's text, keeping its first run's formatting."""
    runs = p.runs
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    runs[0].text = text
    runs[0].font.color.rgb = RED if "<" in text else None


# ---------------------------------------------------------------- document builder
class Report:
    def __init__(self, template):
        self.doc = Document(template)
        self.toc, self.figs, self.tabs = [], [], []
        self.chapter = 0
        self.fig_no = self.tab_no = 0
        self.first_chapter = True
        normal = self.doc.styles["Normal"]
        set_style_font(normal, 12)
        for name, size in (("Heading 1", 16), ("Heading 2", 14)):
            st = self.doc.styles[name]
            set_style_font(st, size, bold=True)
            st.paragraph_format.page_break_before = False

    # ---- title page
    def title_page(self):
        d = DETAILS
        ps = self.doc.paragraphs
        add_bookmark(ps[0], "bm_title")
        set_text(ps[2], f"“{d['title']}”")
        set_text(ps[4], d["course"])
        set_text(ps[6], d["degree"])
        set_text(ps[8], d["branch"])
        student_tpl = ps[11]
        lines = d["students"] + [f"Section: {d['section']}    Year: {d['year']}"]
        for p in ps[12:14]:
            p._p.getparent().remove(p._p)
        set_text(student_tpl, lines[0])
        anchor = student_tpl
        for text in lines[1:]:
            new = copy.deepcopy(student_tpl._p)
            anchor._p.addnext(new)
            from docx.text.paragraph import Paragraph
            anchor = Paragraph(new, student_tpl._parent)
            set_text(anchor, text)
        ps = self.doc.paragraphs
        idx = next(i for i, p in enumerate(ps) if p.text.strip() == "Faculty:")
        set_text(ps[idx + 1], d["faculty"])
        set_text(ps[idx + 2], d["designation"])
        dept = next(i for i, p in enumerate(ps) if p.text.startswith("Department of"))
        set_text(ps[dept], d["department"])
        year_i = next(i for i, p in enumerate(ps) if p.text.strip().startswith("[20"))
        set_text(ps[year_i], d["academic_year"])
        # drop everything after the title page (template instructions)
        body = self.doc.element.body
        el = ps[year_i]._p.getnext()
        while el is not None and el.tag != qn("w:sectPr"):
            nxt = el.getnext()
            body.remove(el)
            el = nxt

    # ---- generic blocks
    def page_break(self):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
        spacing(p, after=0, line=1.0)

    def front_title(self, text, bm):
        p = self.doc.add_paragraph()
        spacing(p, after=18, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        fmt_run(p.add_run(text), 16, bold=True)
        add_bookmark(p, bm)
        return p

    def h1(self, title):
        self.chapter += 1
        self.fig_no = self.tab_no = 0
        p = self.doc.add_paragraph(style="Heading 1")
        spacing(p, before=0, after=6, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        # start each chapter on a new page without an extra break paragraph (which can leave a blank page)
        p.paragraph_format.page_break_before = not self.first_chapter
        self.first_chapter = False
        fmt_run(p.add_run(f"CHAPTER {self.chapter}"), 16, bold=True)
        p.add_run().add_break()
        fmt_run(p.add_run(title.upper()), 16, bold=True)
        bm = f"bm_ch{self.chapter}"
        add_bookmark(p, bm)
        self.toc.append((1, f"{self.chapter}.", title.upper(), bm))
        self.sub = 0
        sp = self.doc.add_paragraph(); spacing(sp, after=6)

    def h2(self, title):
        self.sub += 1
        num = f"{self.chapter}.{self.sub}"
        p = self.doc.add_paragraph(style="Heading 2")
        spacing(p, before=12, after=6, line=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
        fmt_run(p.add_run(f"{num} {title}"), 14, bold=True)
        bm = f"bm_s{num.replace('.', '_')}"
        add_bookmark(p, bm)
        self.toc.append((2, num, title, bm))

    def para(self, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12):
        p = self.doc.add_paragraph()
        spacing(p, align=align)
        rich(p, text, size)
        return p

    def bullets(self, items):
        for text in items:
            p = self.doc.add_paragraph(style="List Paragraph")
            spacing(p, after=3)
            numpr = p._p.get_or_add_pPr().get_or_add_numPr()
            numpr.get_or_add_ilvl().val = 0
            numpr.get_or_add_numId().val = 1
            rich(p, text)

    def equation(self, text):
        p = self.doc.add_paragraph()
        spacing(p, before=6, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        fmt_run(p.add_run(text), 12, italic=True)

    def figure(self, name, caption, interpretation=None, width=15.0):
        self.fig_no += 1
        num = f"{self.chapter}.{self.fig_no}"
        p = self.doc.add_paragraph()
        spacing(p, before=6, after=2, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(FIG / f"{name}.png"), width=Cm(width))
        c = self.doc.add_paragraph()
        spacing(c, after=10, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        fmt_run(c.add_run(f"Figure {num}: "), 11, bold=True)
        fmt_run(c.add_run(caption), 11)
        bm = f"bm_f{num.replace('.', '_')}"
        add_bookmark(c, bm)
        self.figs.append((num, caption, bm))
        if interpretation:
            self.para("**Interpretation:** " + interpretation)

    def table(self, caption, rows, widths, header=True, size=11):
        self.tab_no += 1
        num = f"{self.chapter}.{self.tab_no}"
        c = self.doc.add_paragraph()
        spacing(c, before=6, after=4, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        c.paragraph_format.keep_with_next = True
        fmt_run(c.add_run(f"Table {num}: "), 11, bold=True)
        fmt_run(c.add_run(caption), 11)
        bm = f"bm_t{num.replace('.', '_')}"
        add_bookmark(c, bm)
        self.tabs.append((num, caption, bm))
        t = self.grid(rows, widths, header, size)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return t

    def grid(self, rows, widths, header=True, size=11, bold_rows=()):
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.width = Cm(widths[j])
                p = cell.paragraphs[0]
                spacing(p, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.LEFT)
                if header and i == 0:
                    fmt_run(p.add_run(str(val)), size, bold=True)
                    shade(cell, "D9E2F3")
                elif callable(val):
                    val(p)
                else:
                    rich(p, str(val), size)
                    if i in bold_rows:
                        for r in p.runs:
                            r.bold = True
                mark = OxmlElement("w:rPr")  # empty cells otherwise keep the 12pt default height
                sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); mark.append(sz)
                p._p.get_or_add_pPr().append(mark)
                if len(rows) <= 20 and i < len(rows) - 1:
                    p.paragraph_format.keep_with_next = True  # keep short tables on one page
        return t

    def marker(self):
        return self.doc.add_paragraph()

    # ---- TOC / LOF / LOT, built after the body so every entry is known
    def _move_after(self, marker, table):
        marker._p.addnext(table._tbl)
        marker._p.getparent().remove(marker._p)

    def _page_cell(self, bm, size=11):
        return lambda p: (setattr(p, "alignment", WD_ALIGN_PARAGRAPH.CENTER),
                          add_field(p, f"PAGEREF {bm} \\h", "1", size=size))

    def build_lists(self, toc_marker, lof_marker, lot_marker):
        rows = [["S.No", "Title", "Page"]]
        bold_rows = []
        for title, bm in (("Title Page", "bm_title"), ("Abstract", "bm_abstract"),
                          ("List of Figures", "bm_lof"), ("List of Tables", "bm_lot")):
            rows.append(["", title, self._page_cell(bm, 10.5)])
        for level, num, title, bm in self.toc:
            if level == 1:
                bold_rows.append(len(rows))
                rows.append([num, title, self._page_cell(bm, 10.5)])
            else:
                rows.append(["", f"      {num} {title}", self._page_cell(bm, 10.5)])
        t = self.grid(rows, [1.5, 12.5, 2.0], size=10.5, bold_rows=bold_rows)
        for i in bold_rows:
            for j in (0, 1):
                for r in t.cell(i, j).paragraphs[0].runs:
                    r.bold = True
        self._move_after(toc_marker, t)

        for marker, items, label in ((lof_marker, self.figs, "Figure"), (lot_marker, self.tabs, "Table")):
            rows = [["S.No", f"{label} No.", "Title", "Page"]]
            for k, (num, caption, bm) in enumerate(items, 1):
                rows.append([str(k), f"{label} {num}", caption, self._page_cell(bm)])
            self._move_after(marker, self.grid(rows, [1.3, 2.7, 10.0, 2.0]))

    # ---- sections and page numbers
    def footers(self):
        front, body = self.doc.sections[0], self.doc.sections[1]
        for sec, fmt in ((front, "lowerRoman"), (body, "decimal")):
            pg = OxmlElement("w:pgNumType")
            pg.set(qn("w:fmt"), fmt)
            pg.set(qn("w:start"), "1")
            sec._sectPr.find(qn("w:cols")).addprevious(pg)
        front.different_first_page_header_footer = True
        body.footer.is_linked_to_previous = False
        body.different_first_page_header_footer = False
        for sec in (front, body):
            p = sec.footer.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_field(p, "PAGE", "1", size=11)


# ---------------------------------------------------------------- content
T = lambda n: f"titanic/titanic_{n}"
B = lambda n: f"bank/bank_{n}"
BT = {r["Model"]: r for r in bank["table"]}
BCM = bank["confusion_matrices"]
THR = {r["Threshold"]: r for r in bank["thresholds"]}
RF = BT["Random Forest"]


def build():
    make_diagrams()
    R = Report(BASE / "report_template.docx")
    R.title_page()

    # ---- Abstract
    R.page_break()
    R.front_title("Abstract", "bm_abstract")
    R.para(
        "This report presents two data-science case studies that apply the complete machine-learning workflow to real "
        "classification problems. The first predicts whether a passenger survived the 1912 Titanic disaster from 891 passenger "
        "records; the second predicts whether a customer of a Portuguese bank will subscribe to a term deposit from 41,188 telephone "
        "marketing contacts. In both studies the data was explored with descriptive statistics and visualisations, missing and "
        "unknown values were handled, categorical variables were one-hot encoded and numerical variables scaled, and models were "
        "trained on stratified 80/20 train–test splits. For the Titanic data, a logistic regression model achieved "
        f"{M['Accuracy']:.1%} accuracy and a ROC-AUC of {M['ROC-AUC']:.2f}; its coefficients and p-values showed that sex, passenger "
        "class and age decided survival, and the model was deployed as an interactive Streamlit web application. For the "
        "bank-marketing data, where only 11.3% of customers subscribed, Logistic Regression, Decision Tree, K-Nearest Neighbours and "
        "Random Forest classifiers were compared after excluding the post-call attribute duration; all reached about 90% accuracy, "
        "so precision, recall, F1-score and ROC-AUC were used instead, and Random Forest was selected for its best ROC-AUC "
        f"({RF['ROC-AUC']:.3f}) and its ability to raise recall to {THR[0.2]['Recall']:.0%} with a tuned threshold. Calling the 20% of "
        f"customers it ranks highest would reach {bank['top20_capture']:.0%} of all subscribers, about three times as efficient as "
        "random calling. Together, the studies show how choosing appropriate metrics and interpreting models turns historical data "
        "into reliable predictions and practical, evidence-based decisions.")

    R.page_break(); R.front_title("Table of Contents", "bm_toc"); toc_m = R.marker()
    R.page_break(); R.front_title("List of Figures", "bm_lof"); lof_m = R.marker()
    R.page_break(); R.front_title("List of Tables", "bm_lot"); lot_m = R.marker()
    R.doc.add_section(WD_SECTION.NEW_PAGE)

    # ================================================================ 1. Introduction
    R.h1("Introduction")
    R.h2("Overview of the Case Studies")
    R.para("Classification — predicting which of two or more categories a record belongs to — is one of the most common tasks in "
           "data science. This report applies it to two contrasting problems:")
    R.table("Overview of the two case studies", [
        ["", "Part A – Titanic", "Part B – Bank Marketing"],
        ["Question", "Did a passenger survive?", "Will a customer subscribe to a term deposit?"],
        ["Records", "891 passengers, 11 attributes", "41,188 contacts, 20 attributes"],
        ["Positive class", "38.4% survived", "11.3% subscribed (highly imbalanced)"],
        ["Models", "Logistic Regression", "Logistic Regression, Decision Tree, KNN, Random Forest"],
        ["Focus", "Interpretation of factors and deployment", "Model comparison and business recommendations"],
    ], [3.2, 6.2, 6.6])
    R.h2("Significance and Objectives")
    R.para("Part A turns a historical narrative into measurable evidence, quantifying how gender and class affected the chance of "
           "survival. Part B addresses a practical business problem: telephone campaigns are expensive and most customers decline, so "
           "a model that identifies likely subscribers lets the bank call fewer people and sell more deposits. The objectives are to:")
    R.bullets(["explore both datasets and identify patterns related to the target;",
               "handle missing/unknown values and encode categorical variables;",
               "build a logistic regression model for Titanic survival, evaluate and interpret it, and deploy it with Streamlit;",
               "build and fairly compare four classifiers for bank-marketing prediction using metrics suited to imbalanced data;",
               "recommend the best bank-marketing model and explain how the bank can use it."])
    R.h2("Concepts used in Methodology")
    R.bullets([
        "**Exploratory data analysis:** summary statistics, bar charts, histograms, box plots, correlation heatmaps and pair plots.",
        "**Preprocessing:** median/mode imputation, treating “unknown” as a category, one-hot encoding, standardisation, and "
        "scikit-learn pipelines that prevent data leakage.",
        "**Algorithms:** Logistic Regression (sigmoid of a linear combination), Decision Tree (recursive if–then splits), "
        "K-Nearest Neighbours (majority vote of the k most similar records) and Random Forest (average of many decision trees).",
        "**Evaluation:** confusion matrix, accuracy, precision, recall, F1-score, ROC curve and AUC, and threshold analysis.",
        "**Interpretation and deployment:** odds ratios and p-values, feature importances, and a Streamlit web application."])
    R.h2("Outline of the Results")
    R.para(f"The Titanic logistic regression model reached **{M['Accuracy']:.1%} accuracy** and **{M['ROC-AUC']:.2f} ROC-AUC**, with "
           f"sex, class and age as the significant factors. In the bank study, **Random Forest** achieved the best ROC-AUC "
           f"(**{RF['ROC-AUC']:.3f}**); with a 0.2 threshold it finds {THR[0.2]['Recall']:.0%} of subscribers, and its top-ranked 20% "
           f"of customers contain {bank['top20_capture']:.0%} of all subscribers. Chapters 2–5 present Part A, Chapters 6–9 Part B, "
           "and Chapter 10 concludes. All code and outputs are in the Jupyter Notebook *Data_Science_Case_Studies.ipynb*.")

    # ================================================================ PART A
    R.h1("Part A – Titanic: Problem and Data Exploration")
    R.h2("Problem Statement")
    R.para("On 15 April 1912 the RMS Titanic sank after striking an iceberg; there were not enough lifeboats, and more than 1,500 of "
           "the roughly 2,224 people aboard died. The evacuation followed a “women and children first” policy, and upper-class cabins "
           "were closer to the boat deck. **Problem:** given a passenger's ticket class, sex, age, number of siblings/spouses and "
           "parents/children aboard, fare and port of embarkation, predict whether the passenger **survived (1)** or **did not (0)**, "
           "and identify which factors mattered most.")
    R.h2("Dataset and Summary Statistics")
    R.para(f"*titanic.csv* contains **{len(df)} passengers** and **{df.shape[1]} columns**: PassengerId, Survived (target), Pclass "
           "(1 = upper, 2 = middle, 3 = lower), Name, Sex, Age, SibSp, Parch, Ticket, Fare, Cabin and Embarked (C = Cherbourg, "
           "Q = Queenstown, S = Southampton).")
    desc = df[["Age", "Fare", "SibSp", "Parch", "Survived"]].describe().T
    rows = [["Feature", "Count", "Mean", "Std", "Min", "Median", "Max"]]
    for name, r in desc.iterrows():
        rows.append([name, f"{r['count']:.0f}", f"{r['mean']:.2f}", f"{r['std']:.2f}", f"{r['min']:.2f}",
                     f"{r['50%']:.2f}", f"{r['max']:.2f}"])
    R.table("Titanic – summary statistics of numeric features", rows, [2.6, 2.0, 2.2, 2.2, 2.2, 2.4, 2.4])
    R.para("Passengers were mostly young adults (mean age 29.7). Fare is strongly right-skewed (mean £32.20 vs median £14.45, "
           "maximum £512), most passengers travelled without family, and only **38.4%** survived.")
    R.h2("Missing Values")
    R.para("Cabin is missing for 77.1% of passengers and was dropped. Age is missing for 19.9% (177 rows) and was imputed with "
           "the median; Embarked is missing for only 2 rows and was filled with the most frequent port, Southampton.")
    R.h2("Visual Analysis")
    R.figure(T("02_target_distribution"), "Titanic – distribution of the target variable",
             "549 passengers (61.6%) died and 342 (38.4%) survived. A model that always predicts “died” would be 61.6% accurate, "
             "so precision, recall, F1-score and ROC-AUC are reported alongside accuracy.", width=9)
    R.figure(T("03_survival_by_category"), "Titanic – survival rate by gender, passenger class and port",
             "74% of women survived versus 19% of men. Survival falls from 63% in 1st class to 47% in 2nd and 24% in 3rd. "
             "Cherbourg passengers survived more often (55%) than Southampton passengers (34%), mainly because more of them "
             "travelled 1st class. Women survived more often than men in every class (97%, 92% and 50% vs 37%, 16% and 14%), so "
             "gender and class act independently.", width=16)
    R.figure(T("05_age_hist"), "Titanic – age distribution by survival",
             "Ages peak between 20 and 35. Children under 10 survived far more often (61%) than older passengers (39%), so the "
             "effect of age is concentrated among the very young (overall correlation only −0.08).", width=12)
    R.figure(T("06_fare_box"), "Titanic – fare by class and survival (log scale)",
             "Fares rise sharply with class. Survivors paid a higher median fare (£26) than non-survivors (£10.50); because fare and "
             "class both reflect wealth, their effects overlap.", width=12)
    R.figure(T("07_correlation"), "Titanic – correlation matrix of numeric features",
             "Being male has the strongest correlation with survival (−0.54), followed by Pclass (−0.34) and Fare (+0.26). Pclass and "
             "Fare are correlated (−0.55), as are SibSp and Parch (0.41); no pair exceeds |0.8|, so multicollinearity is not a concern.",
             width=10)
    R.figure(T("08_pairplot"), "Titanic – pair plot of numeric features by survival",
             "No numeric feature separates the classes alone, which is why Sex and Pclass are essential. Survivors cluster at higher "
             "fares, and passengers from very large families were almost all non-survivors.", width=12)
    R.para("**Key findings:** gender is the dominant factor, followed by socio-economic status (class and fare) and age (children "
           "were prioritised); small families fared better than solo travellers or large families.")

    R.h1("Part A – Titanic: Preprocessing and Model")
    R.h2("Project Workflow")
    R.figure("report/workflow", "Workflow followed in both case studies", width=16)
    R.h2("Data Preprocessing")
    R.para("PassengerId, Name, Ticket and Cabin were dropped, leaving Pclass, Sex, Age, SibSp, Parch, Fare and Embarked. The data was "
           f"split 80/20 into {info['n_train']} training and {info['n_test']} test passengers, stratified on *Survived*. The split is "
           "made before imputation so no information leaks from the test set.")
    R.table("Titanic – preprocessing steps", [
        ["Step", "Columns", "Method", "Reason"],
        ["Imputation", "Age, Fare", "Median", "Robust to skew and outliers (median age 28.5)"],
        ["Imputation", "Embarked, Pclass, Sex", "Most frequent", "Only 2 missing values (filled with “S”)"],
        ["Scaling", "Age, SibSp, Parch, Fare", "StandardScaler", "Comparable coefficients, faster convergence"],
        ["Encoding", "Pclass, Sex, Embarked", "One-hot, drop first", "0/1 columns; avoids the dummy-variable trap"],
    ], [2.4, 3.8, 3.2, 6.6])
    R.para("The 7 features become 9 model inputs; the baseline categories are **1st class, female and Cherbourg**. All steps sit in a "
           "scikit-learn *Pipeline*, so training data, test data and app inputs are transformed identically.")
    R.h2("Logistic Regression Model and Training")
    R.equation("P(Survived = 1 | x) = 1 / (1 + e^(−z)),   z = β₀ + β₁x₁ + … + βₖxₖ")
    R.para("z is the log-odds of survival; each coefficient β is the change in log-odds for a one-unit increase in a feature, and "
           "e^β is its **odds ratio**. Coefficients are estimated by maximum likelihood (lbfgs solver, L2 regularisation, C = 1.0). "
           "Training accuracy was **81.0%**, close to the test accuracy, so the model does not overfit.")

    R.h1("Part A – Titanic: Evaluation and Interpretation")
    R.h2("Performance Metrics")
    R.table("Titanic – model performance on the test set", [
        ["Metric", "Value", "Meaning"],
        ["Accuracy", f"{M['Accuracy']:.4f}", "Share of passengers classified correctly"],
        ["Precision", f"{M['Precision']:.4f}", "Of predicted survivors, the share who survived"],
        ["Recall", f"{M['Recall']:.4f}", "Of actual survivors, the share identified"],
        ["F1-score", f"{M['F1-score']:.4f}", "Harmonic mean of precision and recall"],
        ["ROC-AUC", f"{M['ROC-AUC']:.4f}", "Probability of ranking a survivor above a non-survivor"],
    ], [3.0, 2.4, 10.6])
    R.para(f"The model is correct for **{M['Accuracy']:.1%}** of unseen passengers, well above the 61.5% baseline. A precision of "
           f"{M['Precision']:.2f} means about four in five predicted survivors truly survived; recall ({M['Recall']:.2f}) is lower "
           "because survivors who resemble typical victims (adult men, 3rd class) are missed. An AUC of "
           f"{M['ROC-AUC']:.2f} indicates good discrimination.")
    R.figure(T("09_confusion_matrix"), "Titanic – confusion matrix",
             f"{TN + TP} of {TN + FP + FN + TP} test passengers were classified correctly ({TN} TN, {TP} TP); there were {FP} false "
             f"positives and {FN} false negatives, which explains the lower recall.", width=8)
    R.figure(T("10_roc_curve"), "Titanic – ROC curve",
             f"The curve lies well above the random-classifier diagonal (AUC = {M['ROC-AUC']:.3f}); lowering the 0.5 threshold would "
             "raise recall at the cost of more false positives.", width=9)
    R.h2("Coefficients, Odds Ratios and Significance")
    R.figure(T("11_coefficients"), "Titanic – logistic regression coefficients with odds ratios",
             "Being male (OR ≈ 0.08) cut the odds of survival by about 92% relative to a comparable woman; 3rd class (OR ≈ 0.15) cut "
             "them by about 85% relative to 1st class and 2nd class roughly halved them (OR ≈ 0.48). Older age (OR ≈ 0.63 per 14.5 "
             "years) and more siblings/spouses (OR ≈ 0.77) also lowered the odds.", width=13)
    rows = [["Feature", "Coefficient", "Odds ratio", "p-value", "Significant"]]
    for r in sorted(info["significance"], key=lambda r: r["p_value"]):
        rows.append([r["feature"], f"{r['coefficient']:.4f}", f"{r['odds_ratio']:.4f}",
                     "< 0.001" if r["p_value"] < 0.001 else f"{r['p_value']:.4f}", r["significant (p<0.05)"]])
    R.table("Titanic – statistical significance of features (statsmodels, unscaled)", rows, [3.2, 3.2, 3.2, 3.2, 3.2])
    R.para("**Sex, Pclass_3 and Age** are highly significant (p < 0.001) and **Pclass_2 and SibSp** significant (p < 0.05); each extra "
           "year of age multiplies the odds of survival by about 0.96. **Fare, Parch and Embarked** are not significant once class "
           "is known. The model confirms the historical account: being female, travelling in a higher class and being young "
           "determined survival.")

    R.h1("Part A – Titanic: Deployment with Streamlit")
    R.para("The trained pipeline is saved to *model/titanic_logreg.pkl* with joblib. The Streamlit app *app.py* loads it once, "
           "collects a passenger's details from sidebar widgets and returns the survival probability, the predicted outcome and a "
           "chart showing how each feature pushed the log-odds up or down. It runs locally with *streamlit run app.py* and can be "
           "deployed on Streamlit Community Cloud.")
    R.figure(T("13_app_explanation"), "Streamlit app – prediction and explanation for a 30-year-old man in 3rd class",
             "The predicted survival probability is 8.8%; the explanation chart shows that being male (Sex_male) and travelling in "
             "3rd class (Pclass_3) push the prediction down the most. For a 28-year-old woman in 1st class the app returns 91.5%.",
             width=14)

    # ================================================================ PART B
    R.h1("Part B – Bank Marketing: Data Understanding and Preprocessing")
    R.h2("Case Study and Problem Statement")
    R.para("A Portuguese bank runs telephone campaigns to sell term deposits and keeps records of customer characteristics, previous "
           "campaign interactions and outcomes. Different algorithms may give different results, so the bank needs to know which "
           "model predicts subscribers most usefully. **Problem:** develop and compare Logistic Regression, Decision Tree, K-Nearest "
           "Neighbours and Random Forest classifiers that predict whether a customer will subscribe (*y* = yes/no), and recommend the "
           "most appropriate model for the campaign.")
    R.h2("Dataset Description")
    R.para("*bank-additional-full.csv* (UCI Machine Learning Repository, semicolon-delimited) contains **41,188 contacts** from May "
           "2008 to November 2010 with 20 input attributes and the target *y*.")
    R.table("Bank – attribute groups", [
        ["Group", "Attributes", "Type"],
        ["Customer", "age; job, marital, education, default, housing, loan", "1 numerical, 6 categorical"],
        ["Current campaign", "contact, month, day_of_week; duration, campaign", "3 categorical, 2 numerical"],
        ["Previous campaign", "pdays, previous; poutcome", "2 numerical, 1 categorical"],
        ["Economic context", "emp.var.rate, cons.price.idx, cons.conf.idx, euribor3m, nr.employed", "5 numerical"],
        ["Target", "y – subscribed to a term deposit (yes / no)", "Binary"],
    ], [3.2, 9.0, 3.8])
    R.para("Customers are aged 17–98 (mean 40). The number of contacts in the campaign has a median of 2 but reaches 56, and "
           "*pdays* = 999 (never contacted before) for about 96% of customers.")
    R.h2("Data Quality: Unknown Values, Duplicates and Target")
    R.table("Bank – categorical attributes containing “unknown”", [
        ["Attribute", "Unknown entries", "Percent"],
        ["default", "8,597", "20.87%"], ["education", "1,731", "4.20%"], ["housing", "990", "2.40%"],
        ["loan", "990", "2.40%"], ["job", "330", "0.80%"], ["marital", "80", "0.19%"],
    ], [5.0, 5.0, 5.0])
    R.para("There are no empty cells, but missing categorical information is coded as “unknown”. It was **kept as a separate "
           "category**, because deleting those rows would discard over 20% of the data and the absence of information may itself be "
           "informative. **12 duplicate rows** were removed, leaving 41,176 customers.")
    R.h2("Preprocessing and Train–Test Split")
    R.bullets([
        "**duration excluded:** it is the length of the last call, known only after the call (a 0-second call always means *no*); "
        "the UCI documentation states it must be discarded for a realistic pre-contact model.",
        "**Target:** yes → 1, no → 0. **Categorical attributes:** one-hot encoded, including “unknown”. **Numerical attributes:** "
        "standardised — essential for KNN, which is distance-based.",
        f"**Split:** 80:20 stratified — {bank['n_train']:,} training and {bank['n_test']:,} test customers, both with an 11.3% "
        f"subscription rate. The 19 predictors become {bank['n_features_encoded']} encoded inputs.",
        "The same split, preprocessing pipeline and target are used for all four models, ensuring a fair comparison."])

    R.h1("Part B – Bank Marketing: Exploratory Data Analysis")
    R.figure(B("01_target_distribution"), "Bank – distribution of the target variable",
             "Only 4,639 of 41,176 customers (11.27%) subscribed, while 36,537 (88.73%) did not. A model that always predicts “no” "
             "would be 88.7% accurate, so accuracy is misleading and the models must be judged on precision, recall, F1-score and "
             "ROC-AUC for the “yes” class.", width=9)
    R.figure(B("02_rate_by_category"), "Bank – subscription rate by job, month, previous outcome and contact type",
             "Customers whose previous campaign was a success subscribed at 65.1%, versus 14.2% after a failure and 8.8% for new "
             "customers. May, the busiest month, has the lowest rate (6.4%), while March, September, October and December reach "
             "44–51%. Students (31.4%) and retired customers (25.3%) subscribe most, blue-collar workers least (6.9%), and "
             "mobile-phone contacts (14.7%) almost three times as often as landline contacts (5.2%).", width=16)
    R.figure(B("03_numeric_boxplots"), "Bank – numerical attributes by subscription outcome",
             "Subscribers were called when the Euribor rate was much lower (median 1.27% vs 4.86%) and employment was lower "
             "(median 5,099 vs 5,196 thousand). They also needed fewer calls (mean 2.05 vs 2.63); age medians are similar (37 vs "
             "38) but customers aged 60+ subscribe at 39.6%. Correlations confirm this: nr.employed (−0.36), pdays (−0.33) and "
             "euribor3m (−0.31) are the numerical attributes most strongly related to subscription.", width=16)
    R.h2("Three Key Observations")
    R.bullets([
        "**Prior relationship matters most among customer attributes:** customers with a previous successful outcome are about six "
        "times more likely to subscribe (65.1% vs 11.3% overall).",
        "**Timing and economic conditions drive success:** the subscription rate is 24.5% when the 3-month Euribor rate is below 2% "
        "but only 4.8% when it is 4% or more, and the busiest month (May) has the lowest success rate.",
        "**Who and how you call matters:** students, retired customers and mobile-phone contacts subscribe most often, while calling "
        "the same customer more than five times yields only 5.5% success compared with 13.0% for a single call."])

    R.h1("Part B – Bank Marketing: Model Development and Comparison")
    R.h2("Model Configuration")
    R.table("Bank – classification algorithms and settings", [
        ["Model", "Settings", "Reason"],
        ["Logistic Regression", "max_iter = 1000", "Linear, interpretable baseline"],
        ["Decision Tree", "max_depth = 6", "Prevents overfitting (an unrestricted tree: train 0.995, test AUC 0.631)"],
        ["K-Nearest Neighbours", "k = 5", "Predicts from the 5 most similar customers"],
        ["Random Forest", "300 trees, min_samples_leaf = 5", "Averages many trees for stable, non-linear predictions"],
    ], [3.8, 4.6, 7.6])
    R.h2("Performance Comparison Table")
    rows = [["Model", "Accuracy", "Precision", "Recall", "F1-score", "ROC-AUC"]]
    for name, r in BT.items():
        rows.append([name] + [f"{r[k]:.4f}" for k in ["Accuracy", "Precision", "Recall", "F1-score", "ROC-AUC"]])
    R.table("Bank – final performance comparison on the test set (threshold 0.5)", rows, [4.0, 2.4, 2.4, 2.4, 2.4, 2.4])
    R.figure(B("05_model_comparison"), "Bank – comparison of accuracy, precision, recall and F1-score",
             "Accuracy is almost identical for all models (0.896–0.900), so it cannot separate them. Logistic Regression, Decision "
             "Tree and Random Forest have the highest precision (≈0.65); KNN has the highest recall (0.31) and F1-score (0.40) at the "
             "default threshold, but lower precision (0.57). All models miss most subscribers at the 0.5 threshold.", width=15)
    R.h2("Confusion Matrices and Classification Errors")
    R.figure(B("06_confusion_matrices"), "Bank – confusion matrices of the four classifiers",
             f"Of the 928 actual subscribers in the test set, between {BCM['KNN'][1][0]} (KNN) and {BCM['Logistic Regression'][1][0]} "
             "(Logistic Regression) are missed. These false negatives are lost sales and are the dominant error for every model. "
             f"False positives are few — {BCM['Logistic Regression'][0][1]} to {BCM['KNN'][0][1]} — and cost only a wasted call. KNN "
             f"finds the most subscribers ({BCM['KNN'][1][1]}) but with the most false positives; Random Forest finds "
             f"{BCM['Random Forest'][1][1]} with only {BCM['Random Forest'][0][1]} false positives.", width=16.5)
    R.h2("ROC Curves")
    R.figure(B("07_roc_curves"), "Bank – ROC curves of the four classifiers",
             f"Random Forest has the highest ROC-AUC ({RF['ROC-AUC']:.3f}), closely followed by Logistic Regression "
             f"({BT['Logistic Regression']['ROC-AUC']:.3f}) and the Decision Tree ({BT['Decision Tree']['ROC-AUC']:.3f}), while KNN "
             f"is clearly lower ({BT['KNN']['ROC-AUC']:.3f}). Random Forest is therefore the best at ranking customers by their "
             "likelihood of subscribing, independent of any single threshold.", width=10)
    R.h2("Identifying Actual Subscribers")
    R.para("Because the 0.5 threshold is too strict when only 11% of customers subscribe, the Random Forest threshold was varied:")
    rows = [["Threshold", "Customers flagged", "Precision", "Recall", "F1-score"]]
    for r in bank["thresholds"]:
        rows.append([f"{r['Threshold']:.2f}", f"{r['Customers to call']:,}", f"{r['Precision']:.3f}", f"{r['Recall']:.3f}",
                     f"{r['F1-score']:.3f}"])
    R.table("Bank – Random Forest performance at different decision thresholds", rows, [3.0, 3.6, 3.0, 3.0, 3.0])
    R.para(f"Lowering the threshold to 0.2 raises recall from {THR[0.5]['Recall']:.2f} to {THR[0.2]['Recall']:.2f} while precision "
           f"falls only to {THR[0.2]['Precision']:.2f}, and F1 rises to {THR[0.2]['F1-score']:.2f} — better than any model at the "
           f"default threshold. Calling only the 20% of customers ranked highest by Random Forest would reach "
           f"{bank['top20_capture']:.0%} of all subscribers with a {bank['top20_rate']:.0%} conversion rate, versus 11% for random "
           "calling.")
    R.h2("Best-Performing Model")
    R.para(f"**Random Forest** is selected. It has the highest ROC-AUC ({RF['ROC-AUC']:.3f}) and high precision ({RF['Precision']:.2f}), "
           f"and because it ranks customers best, lowering its threshold recovers most subscribers, giving the best F1-score "
           f"({THR[0.2]['F1-score']:.2f}). KNN's higher recall at 0.5 comes from a weaker ranking (AUC {BT['KNN']['ROC-AUC']:.3f}) and "
           "disappears once thresholds are tuned. Accuracy was not used to choose, because all models score about 90%. Its most "
           "important inputs are euribor3m, nr.employed, age, pdays and poutcome = success, consistent with the EDA.")

    R.h1("Part B – Bank Marketing: Interpretation and Recommendations")
    R.h2("Suitability of the Random Forest Model")
    R.bullets(["It **ranks customers best** (highest ROC-AUC), which is exactly how a call list is built.",
               "It **captures non-linear effects** (months, interest-rate levels) and is robust to outliers.",
               "It is **stable and resistant to overfitting** because it averages 300 trees.",
               "Its threshold can be tuned to the business costs: a **false negative is a lost sale**, while a **false positive "
               "is only an extra call**, so recall is worth more than raw accuracy."])
    R.h2("Business Recommendations")
    R.bullets([
        f"**Score and rank customers before each campaign** and call them in order of predicted probability — the top 20% contains "
        f"about {bank['top20_capture']:.0%} of likely subscribers at roughly three times the conversion rate of random calling.",
        "**Use a decision threshold of about 0.2–0.3** to capture 50–60% of subscribers instead of 24%, accepting moderately more "
        "calls because each missed subscriber is lost revenue.",
        "**Target high-potential segments and time campaigns well:** prioritise customers with a previous successful outcome, "
        "students, retired customers and mobile contacts; campaign when interest rates are low; and limit repeated calls to "
        "about three per customer."])
    R.h2("Limitations")
    R.bullets([
        "**Class imbalance:** with only 11% subscribers, all models favour predicting “no” at the default threshold. Results depend "
        "strongly on the threshold, and techniques such as class weighting, SMOTE or cost-sensitive learning were not explored and "
        "could change the model ranking.",
        "**Time period and economic indicators:** the strongest predictors are macro-economic values from the 2008–2010 financial "
        "crisis, so the model may not generalise to different economic conditions and would need retraining."])

    # ================================================================ Conclusion
    R.h1("Conclusion and References")
    R.h2("Conclusion")
    R.para("Both case studies followed the complete data-science workflow — understanding and cleaning data, exploratory analysis, "
           "preprocessing, modelling, evaluation, interpretation and communication of results. In **Part A**, a logistic regression "
           f"model predicted Titanic survival with {M['Accuracy']:.1%} accuracy and {M['ROC-AUC']:.2f} ROC-AUC; its coefficients and "
           "p-values showed that sex, passenger class and age decided survival, and it was deployed as an interactive Streamlit app. "
           "In **Part B**, four classifiers were compared on highly imbalanced bank-marketing data. Accuracy was nearly identical for "
           "all of them, so the evaluation relied on precision, recall, F1-score and ROC-AUC together with the business costs of "
           f"errors. Random Forest ranked customers best (ROC-AUC {RF['ROC-AUC']:.3f}) and, with a tuned threshold, identified about "
           f"{THR[0.2]['Recall']:.0%} of subscribers, making it the recommended model for targeting the bank's campaigns. The main "
           "lesson across both studies is that the right evaluation metric and a clear interpretation of the model matter as much as "
           "the choice of algorithm.")
    R.h2("References")
    refs = ["Kaggle, “Titanic – Machine Learning from Disaster” dataset. https://www.kaggle.com/c/titanic",
            "S. Moro, P. Cortez and P. Rita, “A data-driven approach to predict the success of bank telemarketing,” Decision "
            "Support Systems, vol. 62, pp. 22–31, 2014.",
            "UCI Machine Learning Repository, “Bank Marketing” dataset. https://archive.ics.uci.edu/dataset/222/bank+marketing",
            "F. Pedregosa et al., “Scikit-learn: Machine Learning in Python,” Journal of Machine Learning Research, vol. 12, "
            "pp. 2825–2830, 2011.",
            "L. Breiman, “Random Forests,” Machine Learning, vol. 45, pp. 5–32, 2001.",
            "S. Seabold and J. Perktold, “statsmodels: Econometric and statistical modeling with Python,” Proc. 9th Python in "
            "Science Conference, 2010.",
            "Streamlit documentation, “Deploy your app on Community Cloud.” "
            "https://docs.streamlit.io/streamlit-community-cloud/deploy-your-app"]
    for i, ref in enumerate(refs, 1):
        p = R.para(f"[{i}] {ref}", align=WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.9)

    R.build_lists(toc_m, lof_m, lot_m)
    R.footers()
    R.doc.core_properties.title = DETAILS["title"]
    R.doc.save(OUT)
    print(f"Saved {OUT.name}: {len(R.toc)} headings, {len(R.figs)} figures, {len(R.tabs)} tables")


if __name__ == "__main__":
    build()
